# -*- coding: utf-8 -*-
import json, os
import cv2, numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from quizzes import GROUPS, SESTA

BASE = "https://nuovipercorsi123-sys.github.io/Gottardo-quiz/"
OUT = "/home/claude/build/out"
POSTER = "/home/claude/gottardo-quiz/locandina2.jpg"
mp = json.load(open(f"{OUT}/mappa.json"))
os.makedirs(f"{OUT}/qr", exist_ok=True)

# ---------- QR ----------
enc = cv2.QRCodeEncoder.create()
det = cv2.QRCodeDetector()
def make_qr(url, path):
    m = enc.encode(url)
    m = cv2.copyMakeBorder(m, 4, 4, 4, 4, cv2.BORDER_CONSTANT, value=255)
    big = cv2.resize(m, None, fx=12, fy=12, interpolation=cv2.INTER_NEAREST)
    cv2.imwrite(path, big)
    back = cv2.imread(path, cv2.IMREAD_GRAYSCALE)
    got, _, _ = det.detectAndDecode(back)
    assert got == url, (url, got)

urls = {}
for g in GROUPS:
    for d in g["domande"]:
        key = f'{g["id"]}-{d["letter"]}'
        url = BASE + mp[key] + ".html"
        urls[key] = url
        make_qr(url, f"{OUT}/qr/{key}.png")
print("QR ok:", len(urls))

# ---------- helpers docx ----------
FONT = "Arial"
def set_font(run, size, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size); run.bold = bold; run.italic = italic
    if color: run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

def para(container, text="", size=10, bold=False, italic=False, color=None, align=None, before=0, after=3, keep=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.08
    if keep: pf.keep_with_next = True
    if align == "c": p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text); set_font(r, size, bold, italic, color)
    return p

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor.lstrip("#"))
    tcPr.append(shd)

def cell_borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        v = kw.get(edge)
        e = OxmlElement(f"w:{edge}")
        if v:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(v[0])); e.set(qn("w:color"), v[1]); e.set(qn("w:space"), "0")
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)

def set_widths(table, widths_cm):
    table.autofit = False
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)

def row_height(row, cm, exact=False):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast"); trPr.append(h)

def no_split(row):
    trPr = row._tr.get_or_add_trPr(); e = OxmlElement("w:cantSplit"); trPr.append(e)

def first_para(cell):
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    return p

def rule_lines(doc, n=3, h=0.92):
    t = doc.add_table(rows=n, cols=1); set_widths(t, [17.4])
    for r in t.rows:
        row_height(r, h, exact=True); no_split(r)
        c = r.cells[0]; cell_borders(c, bottom=(4, "999999")); cell_margins(c, 0, 0, 0, 0)
        first_para(c)
    return t

def new_doc():
    doc = Document()
    s = doc.sections[0]
    s.page_width = Cm(21); s.page_height = Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8); s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.4)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return doc

# ---------- fogli alunno ----------
STEPS = [
 "A turno, chi ha il QR apre il proprio quiz e legge ad alta voce domanda e risposte.",
 "Gli altri cercano nel libro e dicono dove si trova la risposta.",
 "Chi ha il telefono sceglie e va avanti. L’esito si vede solo alla fine.",
 "Se qualcosa è sbagliato si ricomincia (le domande cambiano di ordine). Al terzo tentativo trovate un aiuto.",
 "Quando il quiz riesce, scrivete insieme la risposta alla domanda, con parole vostre e con le pagine. Tutti la copiano sul proprio foglio.",
 "Segnate la parola ricevuta nella casella giusta della frase.",
]
INTRO = ("Ieri avete visto la prima parte di Gotthard. Un film è una ricostruzione per lo schermo; "
         "il libro è una ricostruzione storica divulgativa. Oggi, in gruppo, controllate che cosa dicono "
         "le sue pagine sulle persone e sui fatti del Gottardo.")
SPOILER = "Attenzione: il libro racconta anche come finiscono le vicende di Favre ed Escher, che nel film non avete ancora visto."

IMG_CAP = {"E": "Dal libro, pp. 42–45: il Politecnico di Zurigo, la sede del Credito Svizzero e il monumento ad Alfred Escher (p. 45).",
           "F": "Dal libro, p. 49: il trasporto della salma di Favre e la cartolina in suo onore.",
           "A": "Dal libro, p. 57: i precari alloggi degli operai presso Göschenen, disegno di Joseph Nieriker, verso il 1880.",
           "C": "Dal libro, p. 59: cambio dei turni a Göschenen, dipinto di Philipp Fleischer (1886)."}

def sheet(doc, g, d, idx):
    key = f'{g["id"]}-{d["letter"]}'
    # --- pagina 1 ---
    t = doc.add_table(rows=1, cols=2); set_widths(t, [6.2, 11.2])
    c0, c1 = t.rows[0].cells
    for c in (c0, c1): cell_borders(c); cell_margins(c, 0, 0, 0, 120)
    first_para(c0).add_run().add_picture(POSTER, width=Cm(4.4))
    p = first_para(c1); r = p.add_run("Dal film al libro"); set_font(r, 19, True)
    para(c1, "La Ferrovia del Gottardo", 11, False, True, "555555", after=5)
    para(c1, INTRO, 10, after=0)
    para(doc, "", 4, after=2)
    para(doc, "Nome e cognome: ______________________________     Classe: ________     Data: ____________", 10, after=6)
    # banner gruppo
    b = doc.add_table(rows=1, cols=1); set_widths(b, [17.4])
    c = b.rows[0].cells[0]; shade(c, g["color"]); cell_borders(c); cell_margins(c, 90, 90, 160, 160)
    p = first_para(c); r = p.add_run(f'GRUPPO {g["name"].upper()}  ·  {g["pages"]}'); set_font(r, 12, True, color="#FFFFFF")
    p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(f'Tu sei l’alunno n. {"abcde".index(d["letter"]) + 1}: il tuo quiz è la domanda {d["letter"]}) {d["short"]}'); set_font(r, 10.5, False, color="#FFFFFF")
    para(doc, "", 4, after=2)
    # QR + istruzioni
    t2 = doc.add_table(rows=1, cols=2); set_widths(t2, [5.8, 11.6])
    q0, q1 = t2.rows[0].cells
    cell_borders(q0, top=(8, g["color"]), left=(8, g["color"]), bottom=(8, g["color"]), right=(8, g["color"])); cell_margins(q0, 80, 80, 80, 80)
    cell_borders(q1); cell_margins(q1, 0, 0, 200, 0)
    p = first_para(q0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(f"{OUT}/qr/{key}.png", width=Cm(4.7))
    pc = q0.add_paragraph(); pc.alignment = WD_ALIGN_PARAGRAPH.CENTER; pc.paragraph_format.space_after = Pt(0)
    set_font(pc.add_run("Il tuo quiz"), 10, True)
    pu = q0.add_paragraph(); pu.alignment = WD_ALIGN_PARAGRAPH.CENTER; pu.paragraph_format.space_after = Pt(0)
    set_font(pu.add_run(mp[key]), 7.5, False, False, "777777")
    p = first_para(q1); set_font(p.add_run("Come si lavora"), 11, True); p.paragraph_format.space_after = Pt(3)
    for i, s in enumerate(STEPS, 1):
        pp = q1.add_paragraph(); pp.paragraph_format.space_after = Pt(2.5); pp.paragraph_format.left_indent = Cm(0.55); pp.paragraph_format.first_line_indent = Cm(-0.55)
        set_font(pp.add_run(f"{i}.  "), 9.5, True); set_font(pp.add_run(s), 9.5)
    para(doc, "", 4, after=2)
    para(doc, "La frase del gruppo", 12, True, color=g["color"], before=4, after=2, keep=True)
    para(doc, "Ogni quiz riuscito vi dà una parola e il suo numero. Scrivetela nella casella giusta: a lavoro finito, leggete la frase.", 9.5, after=4, keep=True)
    t3 = doc.add_table(rows=1, cols=5); set_widths(t3, [3.48] * 5); t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    row_height(t3.rows[0], 2.6, exact=True)
    for i, c in enumerate(t3.rows[0].cells, 1):
        e = (8, "888888"); cell_borders(c, top=e, left=e, bottom=e, right=e); cell_margins(c, 40, 40, 80, 80)
        p = first_para(c); set_font(p.add_run(str(i)), 8, True, color="#888888")
    para(doc, "", 4, after=4)
    from PIL import Image as _I
    _p = f"{OUT}/img/{g['id']}.jpg"; _im = _I.open(_p); _w = min(17.4, 6.8 * _im.width / _im.height)
    pi = doc.add_paragraph(); pi.alignment = WD_ALIGN_PARAGRAPH.CENTER; pi.paragraph_format.space_after = Pt(2)
    pi.add_run().add_picture(_p, width=Cm(_w))
    para(doc, IMG_CAP[g["id"]], 8.5, False, True, "555555", align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(doc, SPOILER, 9, False, True, "555555", before=2, after=0)
    # --- pagina 2 ---
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    pe = doc.paragraphs[-1]; pe.paragraph_format.space_after = Pt(0)
    para(doc, f'Le risposte del gruppo · {g["name"]}', 15, True, color=g["color"], after=2)
    para(doc, "Scrivete insieme la risposta con parole vostre e indicate le pagine. Tutti scrivono tutto sul proprio foglio.", 9.5, italic=True, color="555555", after=6)
    for dd in g["domande"]:
        para(doc, f'{dd["letter"]})  {dd["sheet"]}', 10.5, True, before=6, after=2, keep=True)
        rule_lines(doc)
    para(doc, f'6)  {SESTA}', 10.5, True, before=8, after=2, keep=True)
    rule_lines(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

doc = new_doc()
doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p) if doc.paragraphs else None
n = 0
for g in GROUPS:
    for i, d in enumerate(g["domande"]):
        sheet(doc, g, d, i); n += 1
# togli l'ultima pagina vuota
last = doc.paragraphs[-1]; last._p.getparent().remove(last._p)
doc.save(f"{OUT}/schede_alunni.docx")
print("fogli:", n)
